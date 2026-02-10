"""Export meeting notes to Word (.docx) and plain text."""

import os
from datetime import datetime

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

from summarizer import MeetingSummary


class DocxExporter:
    """Generate OneNote-compatible Word documents with structured meeting notes."""

    def export_meeting_notes(
        self,
        summary: MeetingSummary,
        transcript: str,
        output_path: str,
        meeting_title: str = None,
        meeting_date: str = None,
    ) -> str:
        os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)

        if meeting_date is None:
            meeting_date = datetime.now().strftime("%Y-%m-%d %H:%M")
        if meeting_title is None:
            meeting_title = f"Meeting Notes — {meeting_date}"

        doc = Document()

        # ---- Title ----
        title = doc.add_heading(meeting_title, level=1)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # ---- Meeting info ----
        doc.add_heading("Meeting Information", level=2)
        doc.add_paragraph(f"Date: {meeting_date}")
        if summary.speakers:
            doc.add_paragraph(f"Speakers: {', '.join(summary.speakers)}")

        # ---- Key Discussion Points ----
        if summary.key_points:
            doc.add_heading("Key Discussion Points", level=2)
            for point in summary.key_points:
                doc.add_paragraph(point, style="List Bullet")

        # ---- Decisions Made ----
        if summary.decisions:
            doc.add_heading("Decisions Made", level=2)
            for decision in summary.decisions:
                doc.add_paragraph(decision, style="List Bullet")

        # ---- Action Items (table) ----
        if summary.action_items:
            doc.add_heading("Action Items", level=2)
            table = doc.add_table(rows=1, cols=4)
            table.style = "Light Grid Accent 1"

            headers = ["Owner", "Task", "Deadline", "Priority"]
            for i, h in enumerate(headers):
                cell = table.rows[0].cells[i]
                cell.text = h
                for run in cell.paragraphs[0].runs:
                    run.bold = True

            for item in summary.action_items:
                row = table.add_row().cells
                row[0].text = item.get("owner", "TBD")
                row[1].text = item.get("task", "")
                row[2].text = item.get("deadline", "TBD")
                row[3].text = item.get("priority", "Medium")

        # ---- Key Insights ----
        if summary.insights:
            doc.add_heading("Key Insights", level=2)
            for insight in summary.insights:
                doc.add_paragraph(insight, style="List Bullet")

        # ---- Follow-up ----
        if summary.follow_up:
            doc.add_heading("Follow-up Items", level=2)
            for item in summary.follow_up:
                doc.add_paragraph(item, style="List Bullet")

        # ---- Full Transcript ----
        doc.add_page_break()
        doc.add_heading("Full Transcript", level=2)
        # Add transcript in smaller font paragraphs to avoid one huge block
        for chunk in transcript.split("\n"):
            p = doc.add_paragraph(chunk)
            for run in p.runs:
                run.font.size = Pt(9)

        doc.save(output_path)
        return os.path.abspath(output_path)

    @staticmethod
    def save_transcript_txt(transcript: str, output_path: str) -> str:
        os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
        with open(output_path, "w", encoding="utf-8") as f:
            f.write(transcript)
        return os.path.abspath(output_path)
